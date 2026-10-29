from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

RESULT_COLUMNS = [
    ("document_key", "Número de documento"),
    ("primary_category", "Categoría principal"),
    ("flags", "Señales"),
    ("duplicate_kind", "Tipo de duplicado"),
    ("messages", "Mensajes"),
    ("erp_document_original", "Documento original ERP"),
    ("dian_document_original", "Documento original DIAN"),
    ("erp_document_normalized", "Documento normalizado ERP"),
    ("dian_document_normalized", "Documento normalizado DIAN"),
    ("erp_third_party_original", "Tercero original ERP"),
    ("dian_third_party_original", "Tercero original DIAN"),
    ("erp_third_party_normalized", "Tercero normalizado ERP"),
    ("dian_third_party_normalized", "Tercero normalizado DIAN"),
    ("erp_date_original", "Fecha original ERP"),
    ("dian_date_original", "Fecha original DIAN"),
    ("erp_date_normalized", "Fecha normalizada ERP"),
    ("dian_date_normalized", "Fecha normalizada DIAN"),
    ("erp_base_original", "Base original ERP"),
    ("dian_base_original", "Base original DIAN"),
    ("erp_base_normalized", "Base normalizada ERP"),
    ("dian_base_normalized", "Base normalizada DIAN"),
    ("base_delta", "Diferencia de base (DIAN - ERP)"),
    ("erp_tax_original", "IVA original ERP"),
    ("dian_tax_original", "IVA original DIAN"),
    ("erp_tax_normalized", "IVA normalizado ERP"),
    ("dian_tax_normalized", "IVA normalizado DIAN"),
    ("tax_delta", "Diferencia de IVA (DIAN - ERP)"),
    ("erp_total_original", "Total original ERP"),
    ("dian_total_original", "Total original DIAN"),
    ("erp_total_normalized", "Total normalizado ERP"),
    ("dian_total_normalized", "Total normalizado DIAN"),
    ("total_delta", "Diferencia de total (DIAN - ERP)"),
    ("erp_rows", "Filas ERP"),
    ("dian_rows", "Filas DIAN"),
]


def build_report(context: dict, results: list[dict], validation_errors: list[dict]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    _write_summary(summary, context)
    _write_result_sheet(
        workbook.create_sheet("Resultados"),
        results,
        "Cada fila es un número de documento. La categoría principal es excluyente; las señales pueden coincidir.",
    )
    _write_result_sheet(
        workbook.create_sheet("Documentos ausentes"),
        [row for row in results if row["primary_category"] in {"MISSING_IN_DIAN", "MISSING_IN_ERP"}],
        "Documentos sin correspondencia válida en la otra fuente.",
    )
    _write_result_sheet(
        workbook.create_sheet("Diferencias de valores"),
        _with_flag(results, "VALUE_MISMATCH"),
        "Incluye documentos cuya categoría principal puede ser otra cuando hay varias diferencias.",
    )
    _write_result_sheet(
        workbook.create_sheet("Posibles duplicados"),
        _with_flag(results, "POSSIBLE_DUPLICATE"),
        "Los duplicados exactos se consolidan para no contar dos veces el mismo documento.",
    )
    _write_result_sheet(
        workbook.create_sheet("Requieren revisión"),
        [row for row in results if row["primary_category"] == "REQUIRES_REVIEW"],
        "La correspondencia o la interpretación no es concluyente. No se eliminaron registros.",
    )
    _write_validation_sheet(workbook.create_sheet("Errores de validación"), validation_errors)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def excel_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value)
    if text[:1] in {"=", "+", "-", "@"} or text.startswith(("\t", "\r")):
        return "'" + text
    return text


def _write_summary(worksheet, context: dict) -> None:
    worksheet.append(["ContaMind AI"])
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet.append(["Informe de conciliación de ventas"])
    worksheet.append([])
    rows = [
        ("Empresa", context.get("company_name")),
        ("Identificación tributaria", context.get("tax_id")),
        ("Periodo desde", context.get("period_start")),
        ("Periodo hasta", context.get("period_end")),
        ("Fecha de ejecución", context.get("executed_at")),
        ("Archivo ERP", context.get("erp_filename")),
        ("Archivo DIAN", context.get("dian_filename")),
        ("Tolerancia monetaria", (context.get("rules") or {}).get("tolerance")),
        ("Clave de correspondencia", (context.get("rules") or {}).get("match_key")),
        ("Campos comparados", ", ".join((context.get("rules") or {}).get("active_fields") or [])),
        ("Aviso", (context.get("rules") or {}).get("disclaimer")),
    ]
    for label, value in rows:
        worksheet.append([label, excel_text(value)])
    worksheet.append([])
    worksheet.append(["Indicador", "Valor"])
    summary = context.get("summary") or {}
    labels = [
        ("erp_rows_read", "Registros leídos en el ERP"),
        ("dian_rows_read", "Registros leídos en la DIAN"),
        ("erp_rows_usable", "Registros con número de documento en el ERP"),
        ("dian_rows_usable", "Registros con número de documento en la DIAN"),
        ("erp_rows_rejected", "Registros del ERP sin número de documento"),
        ("dian_rows_rejected", "Registros de la DIAN sin número de documento"),
        ("matched", "Documentos coincidentes"),
        ("missing_in_dian", "Ausentes en la DIAN"),
        ("missing_in_erp", "Ausentes en el ERP"),
        ("value_mismatch_flags", "Documentos con diferencia de valores"),
        ("date_mismatch_flags", "Documentos con diferencia de fecha"),
        ("third_party_mismatch_flags", "Documentos con diferencia de identificación"),
        ("duplicate_groups", "Grupos con posible duplicado"),
        ("requires_review", "Documentos que requieren revisión"),
        ("processing_errors", "Filas con errores de interpretación"),
        ("result_count", "Resultados de conciliación"),
    ]
    for key, label in labels:
        worksheet.append([label, summary.get(key, 0)])
    skipped = (context.get("rules") or {}).get("skipped_fields") or []
    if skipped:
        worksheet.append([])
        worksheet.append(["Comparaciones no ejecutadas", "Motivo"])
        for item in skipped:
            worksheet.append([excel_text(item.get("field")), excel_text(item.get("reason"))])
    worksheet.column_dimensions["A"].width = 52
    worksheet.column_dimensions["B"].width = 88


def _write_result_sheet(worksheet, results: list[dict], note: str) -> None:
    worksheet.append([excel_text(note)])
    worksheet.append([label for _key, label in RESULT_COLUMNS])
    for cell in worksheet[2]:
        cell.font = Font(bold=True)
    if not results:
        worksheet.append(["No hay registros en esta hoja."])
        return
    for result in results:
        worksheet.append([excel_text(value) for value in _result_values(result)])
    _autosize(worksheet)


def _write_validation_sheet(worksheet, errors: list[dict]) -> None:
    worksheet.append(
        [
            "Una fila puede tener avisos y errores. Un error de interpretación impide usar ese valor en la comparación."
        ]
    )
    headers = ["Fuente", "Fila", "Número de documento", "Severidad", "Campo", "Código", "Mensaje", "Valor original"]
    worksheet.append(headers)
    for cell in worksheet[2]:
        cell.font = Font(bold=True)
    if not errors:
        worksheet.append(["No hay observaciones de validación."])
        return
    for error in errors:
        for issue in error.get("issues") or []:
            field = issue.get("field")
            original = (error.get("original") or {}).get(field) if field else None
            worksheet.append(
                [
                    excel_text(error.get("source")),
                    error.get("row_number"),
                    excel_text(error.get("document_key")),
                    excel_text(issue.get("severity")),
                    excel_text(field),
                    excel_text(issue.get("code")),
                    excel_text(issue.get("message")),
                    excel_text(original),
                ]
            )
    _autosize(worksheet)


def _result_values(result: dict) -> list[object]:
    differences = {item["field"]: item for item in result.get("differences") or []}
    return [
        result.get("document_key"),
        result.get("primary_category"),
        ", ".join(result.get("flags") or []),
        result.get("duplicate_kind"),
        " ".join(result.get("messages") or []),
        _original(result, "erp", "document_number"),
        _original(result, "dian", "document_number"),
        _normalized(result, "erp", "document_number"),
        _normalized(result, "dian", "document_number"),
        _original(result, "erp", "third_party_id"),
        _original(result, "dian", "third_party_id"),
        _normalized(result, "erp", "third_party_id"),
        _normalized(result, "dian", "third_party_id"),
        _original(result, "erp", "issue_date"),
        _original(result, "dian", "issue_date"),
        _normalized(result, "erp", "issue_date"),
        _normalized(result, "dian", "issue_date"),
        _original(result, "erp", "base_amount"),
        _original(result, "dian", "base_amount"),
        _normalized(result, "erp", "base_amount"),
        _normalized(result, "dian", "base_amount"),
        (differences.get("base_amount") or {}).get("delta"),
        _original(result, "erp", "tax_amount"),
        _original(result, "dian", "tax_amount"),
        _normalized(result, "erp", "tax_amount"),
        _normalized(result, "dian", "tax_amount"),
        (differences.get("tax_amount") or {}).get("delta"),
        _original(result, "erp", "total_amount"),
        _original(result, "dian", "total_amount"),
        _normalized(result, "erp", "total_amount"),
        _normalized(result, "dian", "total_amount"),
        (differences.get("total_amount") or {}).get("delta"),
        ", ".join(str(number) for number in result.get("erp_row_numbers") or []),
        ", ".join(str(number) for number in result.get("dian_row_numbers") or []),
    ]


def _original(result: dict, source: str, field: str) -> str | None:
    payload = result.get(f"{source}_original") or {}
    return payload.get(field)


def _normalized(result: dict, source: str, field: str) -> str | None:
    payload = result.get(f"{source}_normalized") or {}
    return payload.get(field)


def _with_flag(results: list[dict], flag: str) -> list[dict]:
    return [row for row in results if flag in (row.get("flags") or []) or row.get("primary_category") == flag]


def _autosize(worksheet) -> None:
    for index in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(index)].width = 28
