import csv
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.core.errors import AppError
from app.modules.reconciliations.engine.normalization import format_decimal

MAX_COLUMNS = 80


@dataclass
class TabularData:
    headers: list[str]
    rows: list[dict[str, str | None]]
    detected_type: str
    encoding: str | None
    warnings: list[str] = field(default_factory=list)


def read_tabular(content: bytes, filename: str, *, max_rows: int) -> TabularData:
    if not content:
        raise AppError("empty_file", "El archivo está vacío.", 422)
    extension = Path(filename or "").suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise AppError(
            "unsupported_file",
            "Solo se admiten archivos CSV y Excel .xlsx.",
            422,
        )
    if content[:2] == b"PK":
        _assert_safe_workbook(content)
        if extension != ".xlsx":
            raise AppError(
                "file_type_mismatch",
                "El contenido es un libro de Excel, pero la extensión no es .xlsx.",
                422,
            )
        return _read_xlsx(content, max_rows)
    if extension != ".csv":
        raise AppError(
            "file_type_mismatch",
            "El contenido no es un archivo Excel .xlsx válido.",
            422,
        )
    return _read_csv(content, max_rows)


def _assert_safe_workbook(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
    except zipfile.BadZipFile as exc:
        raise AppError("invalid_workbook", "El archivo Excel está dañado o no es un .xlsx.", 422) from exc
    if any(name.endswith("vbaProject.bin") or "/vbaProject" in name for name in names):
        raise AppError(
            "macros_not_allowed",
            "El archivo contiene macros. Exporta los valores sin contenido activo y vuelve a cargarlo.",
            422,
        )
    if not any(name.startswith("xl/workbook") for name in names):
        raise AppError("invalid_workbook", "El archivo no tiene la estructura de un libro .xlsx.", 422)


def _read_csv(content: bytes, max_rows: int) -> TabularData:
    last_decode_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError as exc:
            last_decode_error = exc
            continue
        try:
            table = _parse_csv_text(text, max_rows)
        except AppError:
            raise
        except Exception as exc:
            raise AppError(
                "invalid_csv",
                "No fue posible leer el CSV. Revisa el separador y la codificación.",
                422,
            ) from exc
        table.encoding = encoding
        return table
    raise AppError(
        "invalid_encoding",
        "No fue posible identificar la codificación del CSV.",
        422,
    ) from last_decode_error


def _parse_csv_text(text: str, max_rows: int) -> TabularData:
    if text.strip() == "":
        raise AppError("empty_file", "El archivo CSV no tiene contenido.", 422)
    delimiter = _detect_delimiter(text)
    raw_header = [cell.strip() for cell in next(csv.reader(StringIO(text), delimiter=delimiter))]
    _validate_headers(raw_header)
    frame = pd.read_csv(
        StringIO(text),
        dtype=str,
        keep_default_na=False,
        sep=delimiter,
        engine="python",
    )
    if len(frame.columns) != len(raw_header):
        raise AppError("invalid_csv", "No fue posible leer los encabezados del CSV.", 422)
    frame.columns = raw_header
    headers = raw_header
    if len(headers) > MAX_COLUMNS:
        raise AppError("too_many_columns", "El archivo supera el número de columnas admitido.", 422)
    warnings: list[str] = []
    rows: list[dict[str, str | None]] = []
    blank_rows = 0
    for record in frame.to_dict(orient="records"):
        normalized = {
            header: _blank_to_none(str(record[original]).strip())
            for header, original in zip(headers, frame.columns, strict=True)
        }
        if all(value is None for value in normalized.values()):
            blank_rows += 1
            continue
        rows.append(normalized)
        if len(rows) > max_rows:
            raise AppError(
                "file_too_large",
                f"El archivo supera el límite de {max_rows} filas de este proceso.",
                422,
            )
    if not rows:
        raise AppError("empty_file", "El archivo no tiene filas de datos.", 422)
    if blank_rows:
        warnings.append(f"Se omitieron {blank_rows} filas completamente vacías.")
    return TabularData(headers=headers, rows=rows, detected_type="csv", encoding=None, warnings=warnings)


def _read_xlsx(content: bytes, max_rows: int) -> TabularData:
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=True, keep_links=False)
    except Exception as exc:
        raise AppError("invalid_workbook", "No fue posible abrir el archivo Excel.", 422) from exc
    try:
        worksheet = workbook.worksheets[0]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1), None)
        if header_row is None:
            raise AppError("empty_file", "El archivo Excel no tiene encabezados.", 422)
        header_cells = list(header_row)
        while header_cells and _header_value(header_cells[-1].value) is None:
            header_cells.pop()
        headers: list[str] = []
        for cell in header_cells:
            header = _header_value(cell.value)
            if header is None:
                raise AppError("invalid_headers", "Hay encabezados vacíos entre las columnas.", 422)
            headers.append(header)
        _validate_headers(headers)
        if len(headers) > MAX_COLUMNS:
            raise AppError("too_many_columns", "El archivo supera el número de columnas admitido.", 422)
        warnings: list[str] = []
        numeric_columns: set[str] = set()
        rows: list[dict[str, str | None]] = []
        blank_rows = 0
        for cells in worksheet.iter_rows(min_row=2, max_col=len(headers)):
            cell_list = list(cells)
            if any(cell.data_type == "f" for cell in cell_list):
                raise AppError(
                    "formulas_not_allowed",
                    "El archivo contiene fórmulas. Exporta un libro con valores ya calculados, sin fórmulas ni macros.",
                    422,
                )
            normalized = {header: None for header in headers}
            for header, cell in zip(headers, cell_list, strict=False):
                if cell.data_type == "n" and cell.value is not None:
                    numeric_columns.add(header)
                normalized[header] = _excel_cell_to_text(cell.value, getattr(cell, "number_format", None))
            if all(value is None for value in normalized.values()):
                blank_rows += 1
                continue
            rows.append(normalized)
            if len(rows) > max_rows:
                raise AppError(
                    "file_too_large",
                    f"El archivo supera el límite de {max_rows} filas de este proceso.",
                    422,
                )
        if not rows:
            raise AppError("empty_file", "El archivo no tiene filas de datos.", 422)
        if blank_rows:
            warnings.append(f"Se omitieron {blank_rows} filas completamente vacías.")
        if numeric_columns:
            warnings.append(
                "Hay celdas numéricas. Si un número de factura o una identificación se exportó como número, "
                "Excel pudo eliminar los ceros iniciales antes de la carga."
            )
        return TabularData(
            headers=headers,
            rows=rows,
            detected_type="xlsx",
            encoding=None,
            warnings=warnings,
        )
    finally:
        workbook.close()


def _validate_headers(headers: list[str]) -> None:
    if not headers or any(header == "" for header in headers):
        raise AppError("invalid_headers", "El archivo debe tener encabezados con nombre.", 422)
    if len(headers) != len(set(headers)):
        raise AppError("duplicate_headers", "Hay encabezados repetidos. Cada columna debe tener un nombre único.", 422)


def _detect_delimiter(text: str) -> str:
    sample = text.splitlines()[0] if text else ""
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        return dialect.delimiter
    except csv.Error:
        counts = {",": sample.count(","), ";": sample.count(";"), "\t": sample.count("\t")}
        delimiter, count = max(counts.items(), key=lambda item: item[1])
        return delimiter if count else ","


def _header_value(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _blank_to_none(value: str) -> str | None:
    if value == "":
        return None
    return value


def _excel_cell_to_text(value: object, number_format: str | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        decimals = _decimal_places(number_format)
        decimal_value = Decimal(str(value))
        if decimals is not None:
            quantizer = Decimal("1").scaleb(-decimals)
            decimal_value = decimal_value.quantize(quantizer)
        return format_decimal(decimal_value)
    text = str(value).strip()
    return text or None


def _decimal_places(number_format: str | None) -> int | None:
    if not number_format or number_format == "General":
        return None
    if "." not in number_format:
        return 0
    fraction = number_format.split(".")[-1]
    count = 0
    for character in fraction:
        if character in {"0", "#"}:
            count += 1
            continue
        break
    return count
