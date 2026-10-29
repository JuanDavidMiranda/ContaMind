import uuid
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.exc import SQLAlchemyError

from app.db.session import get_engine
from app.modules.identity.models import User

MAPPING = {
    "document_number": "factura",
    "third_party_id": "nit",
    "issue_date": "fecha",
    "base_amount": "base",
    "tax_amount": "iva",
    "total_amount": "total",
}


def csv_file(rows: list[tuple[str, str, str, str, str, str]]) -> bytes:
    lines = ["factura,nit,fecha,base,iva,total"]
    lines.extend(",".join(row) for row in rows)
    return ("\n".join(lines) + "\n").encode()


def create_ready_reconciliation(client, erp_rows, dian_rows, tolerance="0"):
    company = client.post(
        "/api/v1/companies",
        json={"legal_name": "Empresa Sintética SAS", "tax_id": "900000000-1"},
    )
    assert company.status_code == 201
    created = client.post(
        "/api/v1/reconciliations",
        json={
            "company_id": company.json()["id"],
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "tolerance": tolerance,
        },
    )
    assert created.status_code == 201
    reconciliation_id = created.json()["id"]
    for source, rows in (("erp", erp_rows), ("dian", dian_rows)):
        response = client.post(
            f"/api/v1/reconciliations/{reconciliation_id}/files",
            data={"source": source},
            files={"file": (f"{source}.csv", csv_file(rows), "text/csv")},
        )
        assert response.status_code == 200, response.text
    columns = client.get(f"/api/v1/reconciliations/{reconciliation_id}/columns")
    assert columns.status_code == 200
    assert columns.json()["erp"]["suggestions"]["document_number"]["column"] == "factura"
    confirmed = client.put(
        f"/api/v1/reconciliations/{reconciliation_id}/mapping",
        json={"erp": MAPPING, "dian": MAPPING},
    )
    assert confirmed.status_code == 200
    return company.json()["id"], reconciliation_id


def test_health_reports_development_limitations(client):
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok"
    assert body["production_ready"] is False


def test_reconciliation_flow_results_and_report(client):
    _company_id, reconciliation_id = create_ready_reconciliation(
        client,
        [
            ("FV001", "900123456-1", "2026-01-15", "1000.00", "190.00", "1190.00"),
            ("SOLOERP", "900123456-1", "2026-01-15", "10.00", "1.90", "11.90"),
        ],
        [
            ("FV001", "900123456-1", "2026-01-15", "1000.00", "190.00", "1000.00"),
            ("SOLODIAN", "900123456-1", "2026-01-16", "10.00", "1.90", "11.90"),
        ],
    )
    executed = client.post(f"/api/v1/reconciliations/{reconciliation_id}/run")
    assert executed.status_code == 200, executed.text
    summary = executed.json()["summary"]
    assert summary["erp_rows_read"] == 2
    assert summary["dian_rows_read"] == 2
    assert summary["matched"] == 0
    assert summary["value_mismatch_flags"] == 1
    assert summary["missing_in_dian"] == 1
    assert summary["missing_in_erp"] == 1

    mismatches = client.get(
        f"/api/v1/reconciliations/{reconciliation_id}/results",
        params={"category": "VALUE_MISMATCH", "q": "fv001"},
    )
    assert mismatches.status_code == 200
    assert mismatches.json()["total"] == 1
    result_id = mismatches.json()["items"][0]["id"]
    detail = client.get(f"/api/v1/reconciliations/{reconciliation_id}/results/{result_id}")
    assert detail.status_code == 200
    assert detail.json()["primary_category"] == "VALUE_MISMATCH"
    assert any(item["field"] == "total_amount" for item in detail.json()["differences"])

    report = client.get(f"/api/v1/reconciliations/{reconciliation_id}/report")
    assert report.status_code == 200
    workbook = load_workbook(BytesIO(report.content))
    assert "Resumen" in workbook.sheetnames
    assert "Diferencias de valores" in workbook.sheetnames

    history = client.get("/api/v1/reconciliations")
    assert history.status_code == 200
    assert history.json()[0]["status"] == "completed"


def test_other_user_cannot_read_company_or_reconciliation(client):
    company_id, reconciliation_id = create_ready_reconciliation(
        client,
        [("FV001", "900123456-1", "2026-01-15", "10", "1", "11")],
        [("FV001", "900123456-1", "2026-01-15", "10", "1", "11")],
    )
    other_id = uuid.uuid4()
    from sqlalchemy.orm import Session

    with Session(get_engine()) as session:
        session.add(User(id=other_id, full_name="Otra persona", email="otra@example.com", status="active"))
        session.commit()
    headers = {"X-User-Id": str(other_id)}
    assert client.get(f"/api/v1/companies/{company_id}", headers=headers).status_code == 404
    assert client.get(f"/api/v1/reconciliations/{reconciliation_id}", headers=headers).status_code == 404
    report = client.get(f"/api/v1/reconciliations/{reconciliation_id}/report", headers=headers)
    assert report.status_code == 404


def test_database_errors_use_a_stable_response(client, monkeypatch):
    def explode(*_args, **_kwargs):
        raise SQLAlchemyError("conexion rechazada")

    monkeypatch.setattr("app.modules.companies.api.list_companies", explode)
    response = client.get("/api/v1/companies")
    assert response.status_code == 500
    assert response.json()["error"]["code"] == "database_error"
    assert "conexion rechazada" not in response.text


def test_missing_required_column_and_bad_file_are_rejected(client):
    company = client.post("/api/v1/companies", json={"legal_name": "Otra SAS"}).json()
    created = client.post(
        "/api/v1/reconciliations",
        json={"company_id": company["id"], "period_start": "2026-02-01", "period_end": "2026-01-01"},
    )
    assert created.status_code == 422
    created = client.post(
        "/api/v1/reconciliations",
        json={"company_id": company["id"], "period_start": "2026-01-01", "period_end": "2026-01-31"},
    )
    reconciliation_id = created.json()["id"]
    rejected = client.post(
        f"/api/v1/reconciliations/{reconciliation_id}/files",
        data={"source": "erp"},
        files={"file": ("ventas.xlsx", b"%PDF-1.4", "application/octet-stream")},
    )
    assert rejected.status_code == 422
    uploaded = client.post(
        f"/api/v1/reconciliations/{reconciliation_id}/files",
        data={"source": "erp"},
        files={"file": ("erp.csv", b"total\n10\n", "text/csv")},
    )
    assert uploaded.status_code == 200
    client.post(
        f"/api/v1/reconciliations/{reconciliation_id}/files",
        data={"source": "dian"},
        files={"file": ("dian.csv", b"total\n10\n", "text/csv")},
    )
    missing_column = client.put(
        f"/api/v1/reconciliations/{reconciliation_id}/mapping",
        json={
            "erp": {"document_number": "factura", "third_party_id": None, "issue_date": None, "base_amount": None, "tax_amount": None, "total_amount": "total"},
            "dian": {"document_number": "factura", "third_party_id": None, "issue_date": None, "base_amount": None, "tax_amount": None, "total_amount": "total"},
        },
    )
    assert missing_column.status_code == 422
