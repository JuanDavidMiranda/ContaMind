import zipfile
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.errors import AppError
from app.modules.reconciliations.infrastructure.reader import read_tabular
from app.modules.reconciliations.infrastructure.report import build_report, excel_text


def workbook_bytes(rows: list[list[object]], formula: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row_index, column_index, value)
    if formula:
        sheet.cell(2, 2, "=1+1")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_csv_preserves_leading_zeroes_and_reads_semicolon_files():
    content = "factura;total\n00123;1.234,50\n".encode("cp1252")
    table = read_tabular(content, "erp.csv", max_rows=10)
    assert table.rows == [{"factura": "00123", "total": "1.234,50"}]
    assert table.detected_type == "csv"


def test_empty_and_invalid_files_are_rejected():
    with pytest.raises(AppError) as empty_file:
        read_tabular(b"", "ventas.csv", max_rows=10)
    assert empty_file.value.code == "empty_file"
    with pytest.raises(AppError) as header_only:
        read_tabular("factura,total\n".encode(), "ventas.csv", max_rows=10)
    assert header_only.value.code == "empty_file"
    with pytest.raises(AppError) as wrong_type:
        read_tabular(b"%PDF-1.4", "ventas.xlsx", max_rows=10)
    assert wrong_type.value.code == "file_type_mismatch"
    with pytest.raises(AppError) as extension:
        read_tabular(b"factura,total\nFV1,10\n", "ventas.txt", max_rows=10)
    assert extension.value.code == "unsupported_file"


def test_xlsx_rejects_formulas_and_macros():
    with pytest.raises(AppError) as formulas:
        read_tabular(
            workbook_bytes([["factura", "total"], ["FV1", 10]], formula=True),
            "ventas.xlsx",
            max_rows=10,
        )
    assert formulas.value.code == "formulas_not_allowed"

    original = workbook_bytes([["factura"], ["FV1"]])
    packaged = BytesIO()
    with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(packaged, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(AppError) as macros:
        read_tabular(packaged.getvalue(), "ventas.xlsx", max_rows=10)
    assert macros.value.code == "macros_not_allowed"


def test_xlsx_reads_dates_and_decimal_formats_without_executing_content():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "factura"
    sheet["B1"] = "fecha"
    sheet["C1"] = "total"
    sheet["A2"] = "00100"
    sheet["B2"] = "2026-03-01"
    sheet["C2"] = 1190.5
    sheet["C2"].number_format = "#,##0.00"
    buffer = BytesIO()
    workbook.save(buffer)
    table = read_tabular(buffer.getvalue(), "dian.xlsx", max_rows=10)
    assert table.rows[0]["factura"] == "00100"
    assert table.rows[0]["fecha"] == "2026-03-01"
    assert table.rows[0]["total"] in {"1190.5", "1190.50"}


def test_report_protects_formula_injection_and_keeps_sections():
    assert excel_text("=CMD()") == "'=CMD()"
    assert excel_text("+1") == "'+1"
    content = build_report(
        {
            "company_name": "Empresa Sintética SAS",
            "tax_id": "900000000-1",
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "executed_at": "2026-02-01T00:00:00+00:00",
            "erp_filename": "erp.csv",
            "dian_filename": "dian.csv",
            "rules": {
                "tolerance": "0",
                "match_key": "document_number",
                "active_fields": ["total_amount"],
                "skipped_fields": [],
                "disclaimer": "Discrepancia para revisión.",
            },
            "summary": {"matched": 0, "result_count": 1, "erp_rows_read": 1, "dian_rows_read": 1},
        },
        [
            {
                "document_key": "=FV1",
                "primary_category": "VALUE_MISMATCH",
                "flags": ["VALUE_MISMATCH"],
                "duplicate_kind": None,
                "messages": ["El total difiere entre las fuentes."],
                "differences": [
                    {
                        "field": "total_amount",
                        "delta": "-10",
                        "erp_original": "20",
                        "dian_original": "10",
                        "erp_normalized": "20",
                        "dian_normalized": "10",
                    }
                ],
                "erp_original": {"document_number": "=FV1", "total_amount": "20"},
                "dian_original": {"document_number": "=FV1", "total_amount": "10"},
                "erp_normalized": {"document_number": "=FV1", "total_amount": "20"},
                "dian_normalized": {"document_number": "=FV1", "total_amount": "10"},
                "erp_row_numbers": [2],
                "dian_row_numbers": [2],
            }
        ],
        [
            {
                "source": "erp",
                "row_number": 3,
                "document_key": None,
                "original": {"document_number": None},
                "issues": [
                    {
                        "field": "document_number",
                        "severity": "error",
                        "code": "empty_document_number",
                        "message": "El número de documento está vacío.",
                    }
                ],
            }
        ],
    )
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == [
        "Resumen",
        "Resultados",
        "Documentos ausentes",
        "Diferencias de valores",
        "Posibles duplicados",
        "Requieren revisión",
        "Errores de validación",
    ]
    results = workbook["Resultados"]
    assert results["A3"].value == "'=FV1"
    values = workbook["Diferencias de valores"]
    assert values["A3"].value == "'=FV1"
    summary_values = [cell for row in workbook["Resumen"].iter_rows(values_only=True) for cell in row]
    assert "Empresa Sintética SAS" in summary_values
