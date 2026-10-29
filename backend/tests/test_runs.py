import hashlib
import uuid
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db.session import get_engine
from app.modules.reconciliations.infrastructure.orm import ReconciliationEvent, ReconciliationFile, ReconciliationResult, ReconciliationRun
from app.modules.reconciliations.infrastructure.storage import resolve_private_file
from tests.test_api import create_ready_reconciliation


def test_second_execution_preserves_the_first_and_does_not_copy_review(client):
    _company_id, reconciliation_id = create_ready_reconciliation(
        client,
        [("FV001", "900123456-1", "2026-01-15", "1000.00", "190.00", "1190.00"), ("SOLOERP", "900123456-1", "2026-01-16", "10.00", "1.90", "11.90")],
        [("FV001", "900123456-1", "2026-01-15", "1000.00", "190.00", "1000.00"), ("SOLODIAN", "800333444-1", "2026-01-16", "10.00", "1.90", "11.90")],
    )
    first = client.post(f"/api/v1/reconciliations/{reconciliation_id}/run")
    assert first.status_code == 200, first.text
    first_run_id = uuid.UUID(first.json()["current_run_id"])

    with Session(get_engine()) as session:
        original_results = list(session.scalars(select(ReconciliationResult).where(ReconciliationResult.run_id == first_run_id)))
        original_files = list(session.scalars(select(ReconciliationFile).where(ReconciliationFile.run_id == first_run_id)))
        original_events = session.scalar(select(func.count()).select_from(ReconciliationEvent).where(ReconciliationEvent.run_id == first_run_id))
        assert original_results
        assert len(original_files) == 2
        snapshots = {
            row.id: (row.document_key, row.primary_category, row.payload, row.total_delta)
            for row in original_results
        }
        file_snapshots = {row.id: (row.storage_path, row.sha256, resolve_private_file(row.storage_path).read_bytes()) for row in original_files}
        missing = next(row for row in original_results if row.primary_category in {"MISSING_IN_DIAN", "MISSING_IN_ERP"})
        assert missing.total_delta is None
        reviewed = next(row for row in original_results if row.document_key == "FV001")
        reviewed.review_status = "JUSTIFIED"
        session.commit()
        reviewed_id = reviewed.id

    second = client.post(f"/api/v1/reconciliations/{reconciliation_id}/run")
    assert second.status_code == 200, second.text
    assert second.json()["current_run_id"] != first_run_id
    assert second.json()["run_count"] == 2

    with Session(get_engine()) as session:
        for result_id, (document_key, category, payload, delta) in snapshots.items():
            stored = session.get(ReconciliationResult, result_id)
            assert stored is not None
            assert stored.run_id == first_run_id
            assert stored.document_key == document_key
            assert stored.primary_category == category
            assert stored.payload == payload
            assert stored.total_delta == delta
        preserved = session.get(ReconciliationResult, reviewed_id)
        assert preserved.review_status == "JUSTIFIED"
        for file_id, (path, digest, content) in file_snapshots.items():
            stored = session.get(ReconciliationFile, file_id)
            assert stored.storage_path == path
            assert stored.sha256 == digest
            assert resolve_private_file(path).read_bytes() == content
            assert hashlib.sha256(content).hexdigest() == digest
        assert session.scalar(select(func.count()).select_from(ReconciliationEvent).where(ReconciliationEvent.run_id == first_run_id)) == original_events
        second_run = session.get(ReconciliationRun, uuid.UUID(second.json()["current_run_id"]))
        assert second_run.sequence == 2
        second_results = list(session.scalars(select(ReconciliationResult).where(ReconciliationResult.run_id == second_run.id)))
        assert second_results
        assert {row.id for row in second_results}.isdisjoint(snapshots)
        assert all(row.review_status == "PENDING" for row in second_results)
        assert all(row.total_delta is None for row in second_results if row.primary_category in {"MISSING_IN_DIAN", "MISSING_IN_ERP"})

    old_report = client.get(f"/api/v1/reconciliations/{reconciliation_id}/report", params={"run_id": str(first_run_id)})
    new_report = client.get(f"/api/v1/reconciliations/{reconciliation_id}/report")
    assert old_report.status_code == 200
    assert new_report.status_code == 200
    assert old_report.content != b""
    assert "Resumen" in load_workbook(BytesIO(old_report.content)).sheetnames


def test_replacing_a_file_after_completion_keeps_the_finished_copy(client):
    _company_id, reconciliation_id = create_ready_reconciliation(
        client,
        [("FV001", "900123456-1", "2026-01-15", "10", "1", "11")],
        [("FV001", "900123456-1", "2026-01-15", "10", "1", "11")],
    )
    executed = client.post(f"/api/v1/reconciliations/{reconciliation_id}/run")
    assert executed.status_code == 200, executed.text
    first_run_id = uuid.UUID(executed.json()["current_run_id"])
    with Session(get_engine()) as session:
        original = session.scalar(select(ReconciliationFile).where(ReconciliationFile.run_id == first_run_id, ReconciliationFile.source == "erp"))
        original_path = original.storage_path
        original_bytes = resolve_private_file(original_path).read_bytes()

    replaced = client.post(
        f"/api/v1/reconciliations/{reconciliation_id}/files",
        data={"source": "erp"},
        files={"file": ("erp-nuevo.csv", b"factura,nit,fecha,base,iva,total\nFV009,1,2026-01-02,1,0,1\n", "text/csv")},
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["current_run_id"] != first_run_id
    assert replaced.json()["status"] == "files_uploaded"

    with Session(get_engine()) as session:
        preserved = session.scalar(select(ReconciliationFile).where(ReconciliationFile.run_id == first_run_id, ReconciliationFile.source == "erp"))
        assert preserved.storage_path == original_path
        assert resolve_private_file(original_path).read_bytes() == original_bytes
        draft = session.scalar(
            select(ReconciliationFile).where(
                ReconciliationFile.run_id == uuid.UUID(replaced.json()["current_run_id"]),
                ReconciliationFile.source == "erp",
            )
        )
        assert draft.storage_path != original_path
        assert draft.original_filename == "erp-nuevo.csv"
